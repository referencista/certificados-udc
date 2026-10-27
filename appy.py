import io
import json
import re
import unicodedata
import xml.etree.ElementTree as ET
import zipfile
from datetime import datetime, timedelta, timezone

import docx
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
import pandas as pd
import requests
import streamlit as st
import urllib3

# Desactivar advertencias SSL por certificados intermedios institucionales
urllib3.disable_warnings(urllib3.exceptions.InsecureRequestWarning)

# ------------------------------------------------------------------
# CONFIGURACIÓN DE PÁGINA Y ESTILOS UCUENCA
# ------------------------------------------------------------------
st.set_page_config(
    page_title="UCuenca - Generador de Certificados de No Adeudar",
    page_icon="🎓",
    layout="wide",
)

CSS_UCUENCA = """
<style>
    /* Estilo General */
    body {
        font-family: 'Arial', sans-serif;
        background-color: #f4f6f9;
    }
    
    /* Header Institucional */
    .uc-header {
        text-align: center;
        background-color: #13386c !important;
        color: white;
        padding: 20px 30px;
        border-radius: 8px 8px 6px 6px;
        margin-bottom: 25px;
        border-bottom: 5px solid #b81c32 !important;
        box-shadow: 0 4px 10px rgba(0,0,0,0.08);
    }
    .uc-header h1 {
        color: #ffffff !important;
        font-size: 26px !important;
        font-weight: 700 !important;
        margin: 0 !important;
    }
    .uc-header p {
        color: #d1dbe8 !important;
        font-size: 14px !important;
        margin-top: 5px !important;
    }

    /* Botones Institucionales */
    .stButton>button, .stDownloadButton>button {
        background-color: #13386c !important;
        color: #ffffff !important;
        border-radius: 8px !important;
        font-weight: 600 !important;
        border: none !important;
        border-bottom: 4px solid #b81c32 !important;
        padding: 10px 20px !important;
        transition: all 0.2s ease !important;
    }
    
    .stButton>button:hover, .stDownloadButton>button:hover {
        background-color: #0b2347 !important;
        border-bottom: 4px solid #d4223b !important;
        color: #ffffff !important;
    }

    /* Ocultar elementos secundarios predeterminados de Streamlit */
    footer { display: none !important; }
</style>
"""

st.markdown(CSS_UCUENCA, unsafe_allow_html=True)

# Banner de Encabezado Institucional
st.markdown(
    """
    <div class="uc-header">
        <h1>UNIVERSIDAD DE CUENCA</h1>
        <p>Centro de Documentación Regional “Juan Bautista Vázquez” (CDR-JBV) &bull; Certificado de No Adeudar</p>
    </div>
    """,
    unsafe_allow_html=True,
)


# ------------------------------------------------------------------
# LISTA OFICIAL DE REFERENCISTAS ORGANIZADOS POR CAMPUS
# ------------------------------------------------------------------
REFERENCISTAS_POR_CAMPUS = {
    "Campus Central": [
        {"nombre": "DORIS PATRICIA TENESACA CARDENAS", "cargo": "Bibliotecario 2"},
        {"nombre": "ERIKA ELIZABETH IDROVO SALAZAR", "cargo": "Bibliotecario 2"},
        {"nombre": "ERIKA SOFIA PEÑAFIEL VAZQUEZ", "cargo": "Bibliotecario 2"},
        {"nombre": "FRANCISCO TEODORO ASTUDILLO SAQUINAULA", "cargo": "Bibliotecario 2"},
        {"nombre": "JENNY EULALIA PEREZ MEJIA", "cargo": "Bibliotecario 2"},
        {"nombre": "JHOANNA NOEMI MOGOLLÓN GUZMÁN", "cargo": "Bibliotecario 2"},
        {"nombre": "PAOLA DEL ROCIO AMAYA ARCE", "cargo": "Bibliotecario 2"},
        {"nombre": "PATRICIA MARIBEL DUCHI PESÁNTEZ", "cargo": "Bibliotecario 2"},
        {"nombre": "WILMAN GONZALO TANDAZO GUEVARA", "cargo": "Bibliotecario 2"},
        {"nombre": "DIANA ALEXANDRA LEÓN BRAVO", "cargo": "Bibliotecario 2"},
    ],
    "Campus Paraíso": [
        {"nombre": "DANIEL RAMIRO CARRION ROMAN", "cargo": "Bibliotecario 2"},
        {"nombre": "ANDREA", "cargo": "Bibliotecario 2"},
        {"nombre": "XIMENA", "cargo": "Bibliotecario 2"},
    ],
    "Campus Yanuncay": [
        {"nombre": "NUBE DEL ROCIO SALTO MORQUECHO", "cargo": "Bibliotecario 2"},
        {"nombre": "HECTOR BLADIMIR CABRERA RODRIGUEZ", "cargo": "Bibliotecario 2"},
        {"nombre": "LUIS", "cargo": "Bibliotecario 2"},
    ],
    "Campus Centro Histórico": [
        {"nombre": "DIANA", "cargo": "Bibliotecario 2"},
    ],
}

MESES = [
    "enero",
    "febrero",
    "marzo",
    "abril",
    "mayo",
    "junio",
    "julio",
    "agosto",
    "septiembre",
    "octubre",
    "noviembre",
    "diciembre",
]

FACULTADES_MAP = [
    (
        "Facultad de Ciencias Agropecuarias",
        [
            "agropecuaria",
            "agropecuarias",
            "agronomía",
            "agronomia",
            "agronómica",
            "agronomica",
            "veterinaria",
            "medicina veterinaria",
            "zootecnia",
        ],
    ),
    (
        "Facultad de Ciencias de la Hospitalidad",
        [
            "hospitalidad",
            "turismo",
            "gastronomía",
            "gastronomia",
            "hotelería",
            "hoteleria",
        ],
    ),
    (
        "Facultad de Arquitectura y Urbanismo",
        [
            "arquitectura",
            "diseño gráfico",
            "diseño de interiores",
            "diseño interior",
            "urbanismo",
        ],
    ),
    (
        "Facultad de Ciencias Económicas y Administrativas",
        [
            "económica",
            "economía",
            "economica",
            "economia",
            "administración",
            "administracion",
            "contabilidad",
            "auditoría",
            "auditoria",
            "mercadotecnia",
            "finanzas",
            "comercio exterior",
            "empresa",
            "empresas",
        ],
    ),
    (
        "Facultad de Ingeniería",
        [
            "ingeniería civil",
            "ingenieria civil",
            "sistemas",
            "computación",
            "computacion",
            "eléctrica",
            "electrica",
            "electrónica",
            "electronica",
            "telecomunicaciones",
            "industrial",
            "carreteras",
            "software",
        ],
    ),
    (
        "Facultad de Ciencias Médicas",
        [
            "médica",
            "medicina",
            "enfermería",
            "enfermeria",
            "fisioterapia",
            "laboratorio clínico",
            "laboratorio clinico",
            "nutrición",
            "nutricion",
            "salud",
            "fonoaudiología",
            "imagenología",
        ],
    ),
    (
        "Facultad de Ciencias Químicas",
        [
            "química",
            "quimica",
            "bioquímica",
            "bioquimica",
            "farmacia",
            "ingeniería química",
            "ingenieria quimica",
            "ingeniería ambiental",
            "ingenieria ambiental",
            "alimentos",
        ],
    ),
    (
        "Facultad de Filosofía, Letras y Ciencias de la Educación",
        [
            "filosofía",
            "filosofia",
            "educación",
            "educacion",
            "comunicación",
            "comunicacion",
            "idiomas",
            "lengua",
            "historia",
            "pedagogía",
            "pedagogia",
            "literatura",
            "inicial",
            "básica",
            "basica",
        ],
    ),
    (
        "Facultad de Jurisprudencia, Ciencias Políticas y Sociales",
        [
            "jurisprudencia",
            "derecho",
            "trabajo social",
            "orientación familiar",
            "orientacion familiar",
            "política",
            "politica",
            "género",
        ],
    ),
    (
        "Facultad de Artes",
        [
            "artes",
            "música",
            "musica",
            "danza",
            "teatro",
            "artes visuales",
            "diseño teatral",
            "escénicas",
        ],
    ),
    (
        "Facultad de Psicología",
        [
            "psicología",
            "psicologia",
            "psicología clínica",
            "psicologia clinica",
            "psicoeducativa",
            "psicoterapia",
        ],
    ),
]


def normalizar_texto(texto):
    if not texto:
        return ""
    texto = unicodedata.normalize("NFD", str(texto))
    texto = re.sub(r"[\u0300-\u036f]", "", texto)
    return texto.lower()


# ==================================================================
# SISTEMA DE FORMATEO DE CARRERAS (CATÁLOGO CANÓNICO + FALLBACK)
# ==================================================================

CATALOGO_OFICIAL_CARRERAS = {
    # Pregrado
    "pedagogia de las ciencias experimentales": "Pedagogía de las Ciencias Experimentales",
    "pedagogia de la actividad fisica y deporte": "Pedagogía de la Actividad Física y Deporte",
    "pedagogia de los idiomas nacionales y extranjeros": "Pedagogía de los Idiomas Nacionales y Extranjeros",
    "pedagogia de las artes y las humanidades": "Pedagogía de las Artes y las Humanidades",
    "educacion basica": "Educación Básica",
    "educacion inicial": "Educación Inicial",
    "comunicacion": "Comunicación",
    "sociologia": "Sociología",
    "trabajo social": "Trabajo Social",
    "derecho": "Derecho",
    "genero y desarrollo": "Género y Desarrollo",
    "orientacion familiar": "Orientación Familiar",
    "psicologia": "Psicología",
    "psicologia clinica": "Psicología Clínica",
    "psicologia social": "Psicología Social",
    "contabilidad y auditoria": "Contabilidad y Auditoría",
    "administracion de empresas": "Administración de Empresas",
    "economia": "Economía",
    "mercadotecnia": "Mercadotecnia",
    "finanzas": "Finanzas",
    "ingenieria civil": "Ingeniería Civil",
    "ingenieria de sistemas": "Ingeniería de Sistemas",
    "electronica y telecomunicaciones": "Electrónica y Telecomunicaciones",
    "ingenieria electrica": "Ingeniería Eléctrica",
    "ingenieria industrial": "Ingeniería Industrial",
    "ingenieria quimica": "Ingeniería Química",
    "ingenieria ambiental": "Ingeniería Ambiental",
    "bioquimica y farmacia": "Bioquímica y Farmacia",
    "medicina": "Medicina",
    "enfermeria": "Enfermería",
    "fisioterapia": "Fisioterapia",
    "nutricion y dietetica": "Nutrición y Dietética",
    "fonoaudiologia": "Fonoaudiología",
    "imagenologia y radiologia": "Imagenología y Radiología",
    "medicina veterinaria": "Medicina Veterinaria",
    "agronomia": "Agronomía",
    "arquitectura": "Arquitectura",
    "diseno grafico": "Diseño Gráfico",
    "diseno de interiores": "Diseño de Interiores",
    "artes visuales": "Artes Visuales",
    "gastronomia": "Gastronomía",
    "turismo": "Turismo",
    # Posgrado / Maestrías
    "maestria en contabilidad y auditoria": "Maestría en Contabilidad y Auditoría",
    "maestria en gestion publica y buen gobierno": "Maestría en Gestión Pública y Buen Gobierno",
    "maestria en educacion": "Maestría en Educación",
    "maestria en psicologia": "Maestría en Psicología",
}

PALABRAS_TILDADAS_FALLBACK = {
    "pedagogia": "pedagogía",
    "ingenieria": "ingeniería",
    "psicologia": "psicología",
    "sociologia": "sociología",
    "agronomia": "agronomía",
    "gastronomia": "gastronomía",
    "hoteleria": "hotelería",
    "fonoaudiologia": "fonoaudiología",
    "imagenologia": "imagenología",
    "auditoria": "auditoría",
    "maestria": "maestría",
    "biologia": "biología",
    "tecnologia": "tecnología",
    "educacion": "educación",
    "comunicacion": "comunicación",
    "administracion": "administración",
    "gestion": "gestión",
    "investigacion": "investigación",
    "orientacion": "orientación",
    "nutricion": "nutrición",
    "basica": "básica",
    "clinica": "clínica",
    "medica": "médica",
    "fisica": "física",
    "quimica": "química",
    "bioquimica": "bioquímica",
    "politica": "política",
    "publica": "pública",
    "economica": "económica",
    "electronica": "electrónica",
    "electrica": "eléctrica",
    "musica": "música",
    "genero": "género",
    "diseno": "diseño",
    "rediseno": "rediseño",
}

CONECTORES_ESPANOL = {
    "de", "del", "la", "las", "los", "el", "y", "e", "o", "u", "en", "con", "por", "para", "a", "mención", "mencion"
}


def formatear_carrera_espanol(texto):
    if not texto:
        return ""

    texto_limpio = re.sub(r"\s+", " ", str(texto)).strip()
    texto_norm = normalizar_texto(texto_limpio)

    # 1. Búsqueda exacta en el catálogo canónico oficial
    if texto_norm in CATALOGO_OFICIAL_CARRERAS:
        return CATALOGO_OFICIAL_CARRERAS[texto_norm]

    # 2. Formateo de respaldo preservando palabras para corrección en Word
    palabras = texto_limpio.split()
    resultado = []

    for i, palabra in enumerate(palabras):
        prefijo, sufijo = "", ""
        p_work = palabra

        if p_work.startswith("("):
            prefijo, p_work = "(", p_work[1:]
        if p_work.endswith(")"):
            sufijo, p_work = ")", p_work[:-1]

        p_norm = normalizar_texto(p_work)
        base = PALABRAS_TILDADAS_FALLBACK.get(p_norm, p_work.lower())

        if i == 0 or p_norm not in CONECTORES_ESPANOL:
            final = base.capitalize()
        else:
            final = base.lower()

        resultado.append(f"{prefijo}{final}{sufijo}")

    return " ".join(resultado)


# ------------------------------------------------------------------
# EXTRACCIÓN Y LIMPIEZA DE METADATOS VÍA DSPACE 9.3 (MULTI-CAPA)
# ------------------------------------------------------------------
@st.cache_data(ttl=3600, show_spinner=False)
def extraer_metadatos_dspace(url_input):
    url_clean = url_input.strip()

    # Identificar UUID (8-4-4-4-12 hex) o Handle (123456789/xxx)
    match_uuid = re.search(
        r"([a-f0-9]{8}-[a-f0-9]{4}-[a-f0-9]{4}-[a-f0-9]{4}-[a-f0-9]{12})",
        url_clean,
        re.I,
    )
    match_handle = re.search(r"handle/(.+)$", url_clean, re.I) or re.search(
        r"(\d+[\.\d]*/\d+)", url_clean
    )

    uuid_id = match_uuid.group(1) if match_uuid else None
    handle_id = match_handle.group(1).strip() if match_handle else None

    if handle_id and handle_id.endswith("/"):
        handle_id = handle_id[:-1]

    base_apis = [
        "https://rest-dspace.ucuenca.edu.ec/server/api",
        "https://dspace.ucuenca.edu.ec/server/api",
        "https://dspace.ucuenca.edu.ec/api",
    ]

    # Cabeceras navegables completas para superar bloqueos WAF/Fortinet
    headers = {
        "Accept": "application/json, text/html, application/xhtml+xml, */*",
        "Accept-Language": "es-EC,es;q=0.9,en-US;q=0.8,en;q=0.7",
        "Cache-Control": "no-cache",
        "User-Agent": (
            "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36"
            " (KHTML, like Gecko) Chrome/123.0.0.0 Safari/537.36"
        ),
        "Referer": "https://dspace.ucuenca.edu.ec/",
    }

    metadata = {}
    owning_coll_name = ""
    handle_official = (
        f"https://dspace.ucuenca.edu.ec/handle/{handle_id}"
        if handle_id
        else url_clean
    )

    data = None
    exito_extraccion = False

    # --- MÉTODO 1: DSPACE REST API 9.3 EN 2 PASOS ---
    if not uuid_id and handle_id:
        for base in base_apis:
            for pid_param in [
                handle_id,
                f"handle/{handle_id}",
                f"http://hdl.handle.net/{handle_id}",
            ]:
                try:
                    pid_url = f"{base}/pid/find?id={pid_param}"
                    resp_pid = requests.get(
                        pid_url,
                        headers=headers,
                        timeout=8,
                        verify=False,
                        allow_redirects=True,
                    )

                    # Intentar resolver UUID desde la URL redirigida
                    m_u = re.search(
                        r"([a-f0-9]{8}-[a-f0-9]{4}-[a-f0-9]{4}-[a-f0-9]{4}-[a-f0-9]{12})",
                        resp_pid.url,
                        re.I,
                    )
                    if m_u:
                        uuid_id = m_u.group(1)
                        break

                    if resp_pid.status_code == 200:
                        try:
                            d_pid = resp_pid.json()
                            if isinstance(d_pid, dict):
                                uuid_id = d_pid.get("uuid") or d_pid.get("id")
                                if not uuid_id and "_links" in d_pid:
                                    href = (
                                        d_pid["_links"]
                                        .get("self", {})
                                        .get("href", "")
                                    )
                                    m_u2 = re.search(
                                        r"([a-f0-9]{8}-[a-f0-9]{4}-[a-f0-9]{4}-[a-f0-9]{4}-[a-f0-9]{12})",
                                        href,
                                        re.I,
                                    )
                                    if m_u2:
                                        uuid_id = m_u2.group(1)
                                if "metadata" in d_pid:
                                    data = d_pid
                                    break
                        except Exception:
                            pass
                    if uuid_id:
                        break
                except Exception:
                    continue
            if uuid_id or data:
                break

    if uuid_id and not data:
        for base in base_apis:
            try:
                item_url = (
                    f"{base}/core/items/{uuid_id}?embed=owningCollection"
                )
                resp_item = requests.get(
                    item_url, headers=headers, timeout=8, verify=False
                )
                if resp_item.status_code == 200:
                    d_item = resp_item.json()
                    if isinstance(d_item, dict) and "metadata" in d_item:
                        data = d_item
                        break
            except Exception:
                continue

    if data:
        metadata = data.get("metadata", {})
        try:
            owning_coll_name = (
                data.get("_embedded", {})
                .get("owningCollection", {})
                .get("name", "")
                .strip()
            )
        except Exception:
            pass
        if metadata:
            exito_extraccion = True

    # --- MÉTODO 2: PROTOCOLO OAI-PMH (NATIVO DSPACE) ---
    if not exito_extraccion and handle_id:
        try:
            oai_url = (
                f"https://dspace.ucuenca.edu.ec/oai/request?verb=GetRecord&"
                f"identifier=oai:dspace.ucuenca.edu.ec:{handle_id}&metadataPrefix=oai_dc"
            )
            resp_oai = requests.get(
                oai_url, headers=headers, timeout=8, verify=False
            )
            if resp_oai.status_code == 200 and "<OAI-PMH" in resp_oai.text:
                root = ET.fromstring(resp_oai.text)
                ns = {
                    "dc": "http://purl.org/dc/elements/1.1/",
                    "oai_dc": "http://www.openarchives.org/OAI/2.0/oai_dc/",
                }

                meta_oai = {}
                for creator in root.findall(".//dc:creator", ns):
                    if creator.text:
                        meta_oai.setdefault(
                            "dc.contributor.author", []
                        ).append({"value": creator.text.strip()})
                for pub in root.findall(".//dc:publisher", ns):
                    if pub.text:
                        meta_oai.setdefault("dc.publisher", []).append(
                            {"value": pub.text.strip()}
                        )
                for t in root.findall(".//dc:title", ns):
                    if t.text:
                        meta_oai.setdefault("dc.title", []).append(
                            {"value": t.text.strip()}
                        )
                for subj in root.findall(".//dc:subject", ns):
                    if subj.text:
                        meta_oai.setdefault("dc.subject", []).append(
                            {"value": subj.text.strip()}
                        )

                if meta_oai:
                    metadata = meta_oai
                    exito_extraccion = True
        except Exception:
            pass

    # --- MÉTODO 3: WEB SCRAPING HTML SSR ---
    if not exito_extraccion:
        target_urls = [url_clean]
        if handle_id:
            target_urls.append(
                f"https://dspace.ucuenca.edu.ec/handle/{handle_id}?show=full"
            )
            target_urls.append(
                f"https://dspace.ucuenca.edu.ec/handle/{handle_id}"
            )

        for t_url in target_urls:
            try:
                resp_html = requests.get(
                    t_url,
                    headers=headers,
                    timeout=8,
                    verify=False,
                    allow_redirects=True,
                )
                if resp_html.status_code == 200:
                    html_text = resp_html.text
                    meta_dict = {}

                    tags = re.findall(r"<meta\s+([^>]+)>", html_text, re.I)
                    for tag_attr in tags:
                        name_m = re.search(
                            r'(?:name|property|itemprop)=["\']([^"\']+)["\']',
                            tag_attr,
                            re.I,
                        )
                        content_m = re.search(
                            r'content=["\']([^"\']+)["\']', tag_attr, re.I
                        )
                        if name_m and content_m:
                            k = name_m.group(1).strip()
                            v = content_m.group(1).strip()
                            if k and v:
                                meta_dict.setdefault(k, []).append({"value": v})

                    json_ld = re.findall(
                        r'<script\s+type=["\']application/ld\+json["\']\s*>(.*?)</script>',
                        html_text,
                        re.DOTALL | re.I,
                    )
                    for script in json_ld:
                        try:
                            ld_data = json.loads(script.strip())
                            if "author" in ld_data:
                                auths = ld_data["author"]
                                if isinstance(auths, dict):
                                    auths = [auths]
                                for a in auths:
                                    if isinstance(a, dict) and "name" in a:
                                        meta_dict.setdefault(
                                            "dc.contributor.author", []
                                        ).append({"value": a["name"]})
                        except Exception:
                            pass

                    if (
                        meta_dict.get("citation_author")
                        or meta_dict.get("dc.contributor.author")
                        or meta_dict.get("DC.creator")
                    ):
                        metadata = meta_dict
                        exito_extraccion = True
                        break
            except Exception:
                continue

    # --- EXTRAER Y ENSAMBLAR CAMPOS FINALES ---
    def get_all_meta_values(keys_list):
        if isinstance(keys_list, str):
            keys_list = [keys_list]
        res = []
        for key in keys_list:
            items = metadata.get(key, [])
            for item in items:
                v = (
                    item.get("value", "").strip()
                    if isinstance(item, dict)
                    else str(item).strip()
                )
                if v and v not in res:
                    res.append(v)
        return res

    # 1. Autores
    raw_authors = get_all_meta_values([
        "dc.contributor.author",
        "dc.creator",
        "citation_author",
        "DC.creator",
        "author",
        "cris.author.name",
    ])
    autores_formateados = []
    for a in raw_authors:
        a_clean = re.sub(r"\s+", " ", a).strip().upper()
        if a_clean and a_clean not in autores_formateados:
            autores_formateados.append(a_clean)

    # 2. Análisis de dc.publisher (Ej: "Universidad de Cuenca. Facultad de ... Carrera de ...")
    publishers = get_all_meta_values([
        "dc.publisher",
        "thesis.degree.grantor",
        "citation_publisher",
    ])
    pub_facultad, pub_carrera = "", ""
    for pub_text in publishers:
        parts = [p.strip() for p in pub_text.split(".") if p.strip()]
        for part in parts:
            p_norm = normalizar_texto(part)
            if "facultad de" in p_norm and not pub_facultad:
                pub_facultad = part
            elif (
                any(
                    k in p_norm
                    for k in [
                        "carrera de",
                        "maestria en",
                        "maestría en",
                        "programa de",
                        "doctorado en",
                    ]
                )
                and not pub_carrera
            ):
                pub_carrera = part

    # 3. Facultad
    titulos = " ".join(get_all_meta_values(["dc.title", "citation_title"]))
    materias = " ".join(
        get_all_meta_values(["dc.subject", "dc.description.abstract"])
    )
    texto_inferencia = normalizar_texto(f"{titulos} {materias}")

    candidatos_facultad = []
    if pub_facultad:
        candidatos_facultad.append(pub_facultad)
    if owning_coll_name:
        candidatos_facultad.append(owning_coll_name)

    candidatos_facultad.extend(
        get_all_meta_values([
            "thesis.degree.grantor",
            "dc.publisher",
            "dc.department",
            "dc.contributor.department",
            "dc.publisher.department",
            "dc.degree.grantor",
        ])
    )

    facultad_detectada = ""
    for cand in candidatos_facultad:
        cand_norm = normalizar_texto(cand)
        for fac_oficial, kw_list in FACULTADES_MAP:
            fac_norm = normalizar_texto(fac_oficial)
            if fac_norm in cand_norm or any(kw in cand_norm for kw in kw_list):
                facultad_detectada = fac_oficial
                break
        if facultad_detectada:
            break

    if not facultad_detectada:
        for fac_oficial, kw_list in FACULTADES_MAP:
            if any(kw in texto_inferencia for kw in kw_list):
                facultad_detectada = fac_oficial
                break

    # 4. Carrera / Programa (Excluyendo dc.subject para evitar palabras clave)
    candidatos_carrera = []
    if pub_carrera:
        candidatos_carrera.append(pub_carrera)
    if owning_coll_name:
        candidatos_carrera.append(owning_coll_name)

    candidatos_carrera.extend(
        get_all_meta_values([
            "thesis.degree.discipline",
            "thesis.degree.name",
            "dc.degree.discipline",
            "dc.degree.program",
            "dc.publisher.program",
            "dc.department",
            "DC.degree.discipline",
        ])
    )

    carrera_detectada = ""
    for cand in candidatos_carrera:
        cand_clean = re.sub(
            r"^(Universidad de Cuenca[\.\,\-]?\s*|Facultad de [^.]+\.\s*)",
            "",
            cand,
            flags=re.I,
        ).strip()

        if (
            cand_clean
            and "universidad" not in cand_clean.lower()
            and "facultad" not in cand_clean.lower()
        ):
            carrera_detectada = cand_clean
            break

    if not carrera_detectada and candidatos_carrera:
        carrera_detectada = candidatos_carrera[0]

    # 5. Handle Oficial
    uri_items = metadata.get("dc.identifier.uri", [])
    for uri in uri_items:
        val_uri = uri.get("value", "") if isinstance(uri, dict) else str(uri)
        if "handle/" in val_uri or "items/" in val_uri:
            handle_official = val_uri
            break

    return {
        "exito": exito_extraccion,
        "autores": (
            autores_formateados
            if autores_formateados
            else ["APELLIDOS, NOMBRES ESTUDIANTE"]
        ),
        "facultad": (
            facultad_detectada
            if facultad_detectada
            else "Facultad de Filosofía, Letras y Ciencias de la Educación"
        ),
        "carrera": carrera_detectada,
        "handle": handle_official,
    }


# ------------------------------------------------------------------
# GENERADOR DEL DOCUMENTO WORD (.DOCX) - FORMATO INTACTO
# ------------------------------------------------------------------
def crear_documento_word(datos):
    doc = docx.Document()

    # 1. Modificar el proofState existente a 'dirty' para forzar autocorrector en Word
    proof_state = doc.settings.element.find(qn("w:proofState"))
    if proof_state is not None:
        proof_state.set(qn("w:spelling"), "dirty")
        proof_state.set(qn("w:grammar"), "dirty")
    else:
        proof = OxmlElement("w:proofState")
        proof.set(qn("w:spelling"), "dirty")
        proof.set(qn("w:grammar"), "dirty")
        doc.settings.element.append(proof)

    # 2. Configurar o actualizar el idioma a Español (Ecuador)
    style_normal = doc.styles["Normal"]
    rPr = style_normal.element.get_or_add_rPr()
    existing_lang = rPr.find(qn("w:lang"))
    if existing_lang is not None:
        existing_lang.set(qn("w:val"), "es-EC")
    else:
        lang = OxmlElement("w:lang")
        lang.set(qn("w:val"), "es-EC")
        rPr.append(lang)

    # Estilo global Arial
    font_normal = style_normal.font
    font_normal.name = "Arial"
    font_normal.size = Pt(11)

    # Márgenes
    for section in doc.sections:
        section.top_margin = Inches(0.9)
        section.bottom_margin = Inches(0.9)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # ENCABEZADO
    table_header = doc.add_table(rows=1, cols=2)
    table_header.alignment = WD_TABLE_ALIGNMENT.CENTER
    table_header.autofit = False

    cell_left, cell_right = (
        table_header.rows[0].cells[0],
        table_header.rows[0].cells[1],
    )
    cell_left.width = Inches(2.6)
    cell_right.width = Inches(3.9)
    cell_left.vertical_alignment = cell_right.vertical_alignment = (
        WD_ALIGN_VERTICAL.CENTER
    )

    p_logo = cell_left.paragraphs[0]
    p_logo.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p_logo.paragraph_format.space_after = Pt(0)
    run_logo = p_logo.add_run("UCUENCA")
    run_logo.font.name = "Arial"
    run_logo.font.size = Pt(28)
    run_logo.font.bold = True
    run_logo.font.color.rgb = RGBColor(15, 43, 91)

    p_hdr = cell_right.paragraphs[0]
    p_hdr.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_hdr.paragraph_format.line_spacing = 1.15
    p_hdr.paragraph_format.space_after = Pt(0)

    run_hdr = p_hdr.add_run(
        "FORMATO DE NO ADEUDAR MATERIAL BIBLIOGRÁFICO A LA BIBLIOTECA\n"
        "UC-CDRJVB-FOR-020\n"
        "Página 1 de 1"
    )
    run_hdr.font.name = "Arial"
    run_hdr.font.size = Pt(8.5)

    doc.add_paragraph()
    doc.add_paragraph()

    # TÍTULO PRINCIPAL
    p_titulo = doc.add_paragraph()
    p_titulo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_titulo.paragraph_format.space_before = Pt(24)
    p_titulo.paragraph_format.space_after = Pt(24)

    r_tit = p_titulo.add_run("CERTIFICADO DE NO ADEUDAR")
    r_tit.font.name = "Arial"
    r_tit.font.size = Pt(13)
    r_tit.font.bold = True

    # CUERPO DEL CERTIFICADO
    p_cuerpo = doc.add_paragraph()
    p_cuerpo.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_cuerpo.paragraph_format.line_spacing = 1.5
    p_cuerpo.paragraph_format.space_after = Pt(24)

    fac_clean = re.sub(
        r"^(Universidad de Cuenca[\.\,\-]?\s*)", "", datos["facultad"], flags=re.I
    ).strip()
    if not fac_clean.lower().startswith("facultad de") and fac_clean:
        fac_clean = f"Facultad de {fac_clean}"

    carr_clean = str(datos["carrera"]).strip()

    tipo = datos.get("tipo_estudio", "Pregrado")

    r_c1 = p_cuerpo.add_run(
        'El Centro de Documentación Regional "Juan Bautista Vázquez" certifica'
        " que "
    )
    r_c1.font.name = "Arial"

    r_nom = p_cuerpo.add_run(f'{datos["autor"]}')
    r_nom.font.name = "Arial"
    r_nom.font.bold = True

    r_c2 = p_cuerpo.add_run(", portador de la cédula de ciudadanía No. ")
    r_c2.font.name = "Arial"

    r_ced = p_cuerpo.add_run("____________________")
    r_ced.font.name = "Arial"
    r_ced.font.bold = True

    # Redacción limpia según modalidad
    if tipo == "Complexivo":
        texto_cert = (
            f", estudiante de la {fac_clean}, de la {carr_clean}, "
            "modalidad Examen Complexivo, no adeuda ningún bien, ni material bibliográfico en esta dependencia."
        )
    else:
        if carr_clean.lower().startswith(("programa", "doctorado")):
            prefix_carrera = "del"
        else:
            prefix_carrera = "de la carrera de "

        texto_cert = (
            f", estudiante de la {fac_clean} {prefix_carrera} {carr_clean}, "
            "no adeuda ningún bien, ni material bibliográfico en esta dependencia."
        )

    r_c3 = p_cuerpo.add_run(texto_cert)
    r_c3.font.name = "Arial"

    # FECHA (UTC-5 Ecuador)
    p_fecha = doc.add_paragraph()
    p_fecha.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_fecha.paragraph_format.space_after = Pt(24)

    tz_ecuador = timezone(timedelta(hours=-5))
    hoy = datetime.now(tz_ecuador)

    r_fecha = p_fecha.add_run(
        f"Cuenca, {hoy.day} de {MESES[hoy.month - 1]} de {hoy.year}"
    )
    r_fecha.font.name = "Arial"

    # FIRMAS
    doc.add_paragraph().paragraph_format.space_after = Pt(30)

    p_atentamente = doc.add_paragraph()
    p_atentamente.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_at = p_atentamente.add_run(
        "Atentamente,\n\n\n\n\n\n________________________________________"
    )
    r_at.font.name = "Arial"

    p_firma = doc.add_paragraph()
    p_firma.alignment = WD_ALIGN_PARAGRAPH.CENTER

    r_f1 = p_firma.add_run(f'\n{datos["ref_nombre"]}\n')
    r_f1.font.name = "Arial"
    r_f1.font.bold = True

    r_f2 = p_firma.add_run(
        f'{datos["ref_cargo"]}\nCDR "Juan Bautista Vázquez"'
    )
    r_f2.font.name = "Arial"

    # ENLACE DSPACE
    for _ in range(2):
        doc.add_paragraph()

    p_link = doc.add_paragraph()
    r_l1 = p_link.add_run("Link: ")
    r_l1.font.name = "Arial"
    r_l1.font.bold = True

    valor_link = (
        "Examen Complexivo" if tipo == "Complexivo" else datos.get("handle", "")
    )
    r_h = p_link.add_run(valor_link)
    r_h.font.name = "Arial"
    r_h.font.underline = True
    r_h.font.color.rgb = RGBColor(0, 51, 153)

    # VERSIÓN
    doc.add_paragraph()

    p_ver = doc.add_paragraph()
    p_ver.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r_v = p_ver.add_run("Version: 1.0")
    r_v.font.name = "Arial"
    r_v.font.size = Pt(9.5)

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer


# ------------------------------------------------------------------
# INTERFAZ PRINCIPAL EN STREAMLIT
# NUEVO BLOQUE MULTI-CAMPUS
# ------------------------------------------------------------------
# NUEVO BLOQUE MULTI-CAMPUS:
col_campus, col_ref, col_tipo = st.columns([2, 2, 2])

with col_campus:
    campus_sel = st.selectbox(
        "Campus / Sede:",
        list(REFERENCISTAS_POR_CAMPUS.keys()),
    )

with col_ref:
    lista_ref_campus = REFERENCISTAS_POR_CAMPUS[campus_sel]
    nombres_ref = sorted([r["nombre"] for r in lista_ref_campus])
    referencista_sel = st.selectbox("Referencista que firma:", nombres_ref)

with col_tipo:
    tipo_estudio = st.selectbox(
        "Tipo de Titulación / Proceso:",
        ["Pregrado", "Maestría", "Doctorado", "Complexivo"],
    )

ref_info = next(
    item for item in REFERENCISTAS_POR_CAMPUS[campus_sel] if item["nombre"] == referencista_sel
)

st.markdown("---")

# ------------------------------------------------------------------
# CASO A: PREGRADO, MAESTRÍA O DOCTORADO (DSPACE)
# ------------------------------------------------------------------
if tipo_estudio != "Complexivo":
    st.markdown("#### 1. Parámetros de la Consulta DSpace")

    if "handle_input" not in st.session_state:
        st.session_state["handle_input"] = ""

    def limpiar_busqueda():
        st.session_state["handle_input"] = ""
        st.session_state.pop("datos_cargados", None)
        st.session_state.pop("carrera_input", None)

    url_input = st.text_input(
        "Handle o URL de DSpace:",
        placeholder=(
            "Ej: https://dspace.ucuenca.edu.ec/handle/123456789/49693 o"
            " 123456789/49693"
        ),
        key="handle_input",
    )

    col_btn1, col_btn2 = st.columns([3, 1])

    with col_btn1:
        btn_procesar = st.button(
            "🔍 Consultar DSpace y Generar Documento(s)",
            use_container_width=True,
        )

    with col_btn2:
        btn_limpiar = st.button(
            "🧹 Nueva Búsqueda",
            use_container_width=True,
            on_click=limpiar_busqueda,
        )

    if btn_procesar or "datos_cargados" in st.session_state:
        if btn_procesar:
            if not url_input:
                st.warning("⚠️ Por favor ingresa el Handle de DSpace.")
                st.stop()
            st.session_state.pop("datos_cargados", None)

            with st.spinner(
                "Conectando con el repositorio DSpace de la UCuenca..."
            ):
                meta = extraer_metadatos_dspace(url_input)
                st.session_state["datos_cargados"] = meta
                st.session_state["carrera_input"] = formatear_carrera_espanol(
                    meta["carrera"]
                )

        meta = st.session_state["datos_cargados"]

        if not meta.get("exito", False):
            st.warning(
                "⚠️ **Aviso de conexión:** No se pudieron extraer los datos"
                " automáticamente desde el servidor DSpace de la UCuenca"
                " (posible bloqueo de firewall/WAF hacia la nube de Streamlit)."
                " **Puede ingresar o corregir los datos manualmente a"
                " continuación.**"
            )

        st.markdown("---")
        st.markdown("#### 2. Validación de Metadatos y Estudiantes")

        def auto_formatear_carrera_cb():
            st.session_state["carrera_input"] = formatear_carrera_espanol(
                st.session_state["carrera_input"]
            )

        if "carrera_input" not in st.session_state:
            st.session_state["carrera_input"] = formatear_carrera_espanol(
                meta["carrera"]
            )

        col_f, col_c = st.columns(2)
        with col_f:
            facultad_final = st.text_input(
                "Facultad Detectada:", value=meta["facultad"]
            )

        with col_c:
            carrera_final = st.text_input(
                "Carrera / Programa Detectado:",
                key="carrera_input",
                on_change=auto_formatear_carrera_cb,
            )

        st.markdown(
            f"**Estudiantes / Autores encontrados ({len(meta['autores'])})**"
        )

        certificados_generados = []

        for idx, autor_nombre in enumerate(meta["autores"], start=1):
            with st.expander(
                f"🎓 Estudiante #{idx}: {autor_nombre}", expanded=True
            ):
                nom_est = st.text_input(
                    f"Nombre Estudiante #{idx}:",
                    value=autor_nombre,
                    key=f"nom_{idx}",
                )

                payload = {
                    "autor": nom_est.strip().upper(),
                    "facultad": facultad_final.strip(),
                    "carrera": carrera_final.strip(),
                    "tipo_estudio": tipo_estudio,
                    "handle": meta["handle"],
                    "ref_nombre": ref_info["nombre"],
                    "ref_cargo": ref_info["cargo"],
                }
                certificados_generados.append(payload)

                buf = crear_documento_word(payload)
                st.download_button(
                    label=f"📄 Descargar Certificado Word - Estudiante {idx}",
                    data=buf,
                    file_name=f"Certificado_{nom_est.replace(' ', '_')}.docx",
                    mime=(
                        "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
                    ),
                    key=f"btn_dl_{idx}",
                )

        if len(certificados_generados) > 1:
            st.markdown("---")
            zip_buffer = io.BytesIO()
            with zipfile.ZipFile(zip_buffer, "w", zipfile.ZIP_DEFLATED) as zf:
                for c_data in certificados_generados:
                    doc_buf = crear_documento_word(c_data)
                    zf.writestr(
                        f"Certificado_{c_data['autor'].replace(' ', '_')}.docx",
                        doc_buf.getvalue(),
                    )

            zip_buffer.seek(0)
            st.download_button(
                label=(
                    "📦 Descargar TODOS los Certificados en Un Archivo ZIP"
                    f" ({len(certificados_generados)} archivos)"
                ),
                data=zip_buffer,
                file_name="Certificados_No_Adeudar_UCuenca.zip",
                mime="application/zip",
            )

# ------------------------------------------------------------------
# CASO B: EXAMEN COMPLEXIVO
# ------------------------------------------------------------------
else:
    st.markdown("#### 1. Parámetros de Examen Complexivo")

    modo_complexivo = st.radio(
        "Seleccione la modalidad de generación:",
        [
            "👤 Ingreso Manual (1 Estudiante)",
            "📂 Carga Masiva desde Excel (Listado)",
        ],
        horizontal=True,
    )

    st.markdown("---")

    if "Ingreso Manual" in modo_complexivo:
        st.info(
            "ℹ️ Ingrese los datos del estudiante de examen complexivo para"
            " generar su certificado individual."
        )

        col_m1, col_m2 = st.columns(2)
        with col_m1:
            nom_comp = st.text_input(
                "Nombres y Apellidos Completos:",
                value="ZHIZHPON GUACHICHULLCA BYRON GEOVANNY",
            )
            carr_comp = st.text_input(
                "Carrera o Programa:",
                value="PEDAGOGIA DE LAS ARTES Y LAS HUMANIDADES",
            )
        with col_m2:
            fac_comp = st.text_input(
                "Facultad:",
                value="FILOSOFIA, LETRAS Y CIENCIAS DE LA EDUCACION",
            )

        payload_manual = {
            "autor": nom_comp.strip().upper(),
            "facultad": fac_comp.strip(),
            "carrera": formatear_carrera_espanol(carr_comp),
            "tipo_estudio": "Complexivo",
            "handle": "",
            "ref_nombre": ref_info["nombre"],
            "ref_cargo": ref_info["cargo"],
        }

        buf_manual = crear_documento_word(payload_manual)

        st.download_button(
            label="📄 Descargar Certificado Word (.docx)",
            data=buf_manual,
            file_name=(
                f"Certificado_Complexivo_{nom_comp.replace(' ', '_')}.docx"
            ),
            mime=(
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
            ),
            key="btn_complexivo_manual",
        )

    else:
        st.info(
            "📂 Cargue el archivo Excel con el listado de Exámenes Complexivos"
            " para generar los certificados de los estudiantes autorizados."
        )

        archivo_excel = st.file_uploader(
            "Seleccione el archivo Excel (.xlsx, .xls):", type=["xlsx", "xls"]
        )

        if archivo_excel is not None:
            try:
                df_raw = pd.read_excel(archivo_excel, header=None)

                header_idx = 0
                for idx, row in df_raw.iterrows():
                    row_text = " ".join(
                        [str(cell).lower() for cell in row.values if pd.notna(cell)]
                    )
                    if (
                        "facultad" in row_text
                        or "apellidos" in row_text
                        or "nombres" in row_text
                    ):
                        header_idx = idx
                        break

                df = pd.read_excel(archivo_excel, header=header_idx)

                def norm_col(texto):
                    txt = (
                        unicodedata.normalize("NFD", str(texto))
                        .encode("ascii", "ignore")
                        .decode("utf-8")
                    )
                    return txt.strip().lower()

                cols_dict = {norm_col(c): c for c in df.columns}

                def encontrar_columna(palabras_clave):
                    for kw in palabras_clave:
                        for norm_c, orig_c in cols_dict.items():
                            if kw in norm_c:
                                return orig_c
                    return None

                col_fac = encontrar_columna(["facultad"])
                col_carr = encontrar_columna(
                    ["carrera o maestria", "carrera", "programa"]
                )
                col_ape = encontrar_columna(["apellidos", "apellido"])
                col_nom = encontrar_columna(["nombres", "nombre"])
                col_aut = encontrar_columna(["autorizacion", "autoriz"])

                faltantes = []
                if not col_fac:
                    faltantes.append("Facultad")
                if not col_carr:
                    faltantes.append("Carrera o Maestría")
                if not col_ape:
                    faltantes.append("Apellidos")
                if not col_nom:
                    faltantes.append("Nombres")
                if not col_aut:
                    faltantes.append("Autorización")

                if faltantes:
                    st.error(
                        "❌ No se pudieron identificar las siguientes columnas:"
                        f" **{', '.join(faltantes)}**."
                    )
                    st.warning(
                        "📋 **Columnas detectadas en la fila"
                        f" {header_idx + 1} de tu Excel:**\n`{list(df.columns)}`"
                    )
                else:
                    df = df.dropna(how="all")
                    df_filtrado = df[
                        df[col_aut].astype(str).str.strip().str.upper()
                        == "AUTORIZADO"
                    ].copy()
                    cant_autorizados = len(df_filtrado)

                    if cant_autorizados == 0:
                        st.warning(
                            "⚠️ Se leyó el archivo correctamente (encabezado en"
                            f" fila {header_idx + 1}), pero no se encontraron"
                            " filas con el valor 'AUTORIZADO' en la columna"
                            f" '{col_aut}'."
                        )
                    else:
                        st.success(
                            f"✅ Se encontraron **{cant_autorizados}**"
                            " estudiantes con estado **AUTORIZADO** (encabezado"
                            f" detectado en la fila {header_idx + 1})."
                        )

                        st.dataframe(
                            df_filtrado[
                                [col_fac, col_carr, col_ape, col_nom, col_aut]
                            ],
                            use_container_width=True,
                        )

                        if st.button(
                            "⚡ Generar Todos los Certificados (.ZIP)",
                            use_container_width=True,
                        ):
                            with st.spinner(
                                "Procesando certificados y empaquetando en"
                                " archivo ZIP..."
                            ):
                                zip_buffer = io.BytesIO()

                                with zipfile.ZipFile(
                                    zip_buffer, "w", zipfile.ZIP_DEFLATED
                                ) as zf:
                                    for _, row in df_filtrado.iterrows():
                                        nombre_completo = (
                                            f"{row[col_ape]}"
                                            f" {row[col_nom]}".strip().upper()
                                        )
                                        facultad_val = str(
                                            row[col_fac]
                                        ).strip()
                                        carrera_val = formatear_carrera_espanol(
                                            str(row[col_carr]).strip()
                                        )

                                        payload_excel = {
                                            "autor": nombre_completo,
                                            "facultad": facultad_val,
                                            "carrera": carrera_val,
                                            "tipo_estudio": "Complexivo",
                                            "handle": "",
                                            "ref_nombre": ref_info["nombre"],
                                            "ref_cargo": ref_info["cargo"],
                                        }

                                        doc_buf = crear_documento_word(
                                            payload_excel
                                        )
                                        zf.writestr(
                                            "Certificado_"
                                            f"{nombre_completo.replace(' ', '_')}.docx",
                                            doc_buf.getvalue(),
                                        )

                                zip_buffer.seek(0)
                                st.download_button(
                                    label=(
                                        "📦 Descargar Certificados Complexivos"
                                        f" (.ZIP) ({cant_autorizados} archivos)"
                                    ),
                                    data=zip_buffer,
                                    file_name=(
                                        "Certificados_Complexivo_UCuenca.zip"
                                    ),
                                    mime="application/zip",
                                    key="btn_dl_complexivo_zip",
                                )
            except Exception as e:
                st.error(
                    f"❌ Ocurrió un error al procesar el archivo Excel: {e}"
                )
